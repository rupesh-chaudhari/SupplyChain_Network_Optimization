#!/usr/bin/env python3
"""
Supply Chain Network Optimisation
Objective : Profit Maximisation
Model     : MILP (scipy HiGHS branch-and-bound)

Input workbook  — 8 sheets:
  SKU_Parameters, Demand, Production_Lines, Available_Hours,
  Handling_Costs, Safety_Stock, Cost_Parameters, Network

Output workbook — sheets:
  Dashboard, Profit Bridge, Production Schedule, Network Flow,
  Lane Utilisation, Prebuild Plan, Inventory Position,
  Demand Fulfilment, Capacity Utilisation, Line Eligibility Matrix
"""

import sys, os, warnings, time, math
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from scipy.optimize import milp, LinearConstraint, Bounds
import scipy.sparse as sp
warnings.filterwarnings("ignore")

C = dict(
    hd="1F3864", hm="2E75B6", hl="9DC3E6",
    acc="ED7D31", accl="FCE4D6",
    grn="375623", grnl="E2EFDA", grnm="70AD47",
    red="C00000", redl="FFE0E0",
    yel="7F6000", yell="FFFF99",
    pur="5C2D91", purl="E8D5F5",
    teal="1F7979", teall="D6EDED",
    ora="C55A11", tbl="004B6B",
    wht="FFFFFF", gyl="F2F2F2", blk="000000",
)

def F(h):    return PatternFill("solid", fgColor=h)
def B(s=10, c="000000"): return Font(name="Calibri", bold=True,  size=s, color=c)
def R(s=10, c="000000"): return Font(name="Calibri", bold=False, size=s, color=c)
def CA(w=False): return Alignment(horizontal="center", vertical="center", wrap_text=w)
def LA(w=False): return Alignment(horizontal="left",   vertical="center", wrap_text=w)
def RA():        return Alignment(horizontal="right",  vertical="center")
def TB():
    t = Side(style="thin", color="BFBFBF")
    return Border(left=t, right=t, top=t, bottom=t)
def KB():
    t = Side(style="medium", color="2E75B6")
    return Border(left=t, right=t, top=t, bottom=t)

FMT_INR = "₹#,##0"
FMT_IN2 = "₹#,##0.00"
FMT_NUM = "#,##0"
FMT_DEC = "#,##0.0"
FMT_PCT = "0.0"


# ── Configuration ─────────────────────────────────────────────────────────────
# Edit INPUT_FILE to point to your input workbook.
# OUTPUT_FILE: leave None to auto-name as <input>_output.xlsx beside the input.

INPUT_FILE  = "sample_input.xlsx"
OUTPUT_FILE = None


# ── Data Loader ───────────────────────────────────────────────────────────────
class DataLoader:

    REQUIRED = {
        "SKU_Parameters":  ["SKU", "Revenue_per_case_INR"],
        "Demand":          ["Customer", "SKU"],
        "Production_Lines":["Plant", "Line", "SKU", "Month",
                            "Line_SKU_active", "Batch_size_cases",
                            "Run_rate_cases_per_hr", "Min_run_hrs",
                            "CIP_hrs", "Production_cost_per_case_INR",
                            "RMPM_Cost_per_case_INR", "Changeover_cost_INR"],
        "Available_Hours": ["Plant", "Line"],
        "Handling_Costs":  ["Node", "SKU",
                            "Handling_cost_per_case_INR",
                            "Inventory_holding_cost_per_case_per_month_INR"],
        "Safety_Stock":    ["Node", "SKU", "Safety_stock_days"],
        "Cost_Parameters": ["Parameter", "Value"],
        "Network":         ["Source", "Destination", "SKU", "Month",
                            "Transport_cost_per_case_INR", "Lane_open"],
    }

    def __init__(self, path):
        print(f"\n  Loading: {path}")
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except FileNotFoundError:
            raise SystemExit(f"\n  ERROR: File not found: {path}")
        except Exception as e:
            raise SystemExit(f"\n  ERROR: Cannot open workbook: {e}")

        available_sheets = wb.sheetnames
        print(f"   Sheets found: {available_sheets}")

        missing = [s for s in self.REQUIRED if s not in available_sheets]
        if missing:
            raise SystemExit(
                f"\n  ERROR: Missing required sheets: {missing}"
            )

        def sheet_to_df(ws_name):
            ws = wb[ws_name]
            rows = list(ws.iter_rows(values_only=True))
            header_row = None
            for i, row in enumerate(rows):
                non_blank = [c for c in row if c is not None and str(c).strip() != ""]
                if len(non_blank) >= 2:
                    header_row = i
                    break
            if header_row is None:
                raise SystemExit(
                    f"\n  ERROR: Sheet '{ws_name}' is empty or has no header row."
                )
            headers = [str(c).strip() if c is not None else f"_col{ci}"
                       for ci, c in enumerate(rows[header_row])]
            data_rows = []
            for row in rows[header_row + 1:]:
                if all(c is None or str(c).strip() == "" for c in row):
                    continue
                first = str(row[0]).strip() if row[0] is not None else ""
                if first.startswith("↑") or first.startswith("►"):
                    continue
                data_rows.append(list(row))
            df = pd.DataFrame(data_rows, columns=headers)
            return df.dropna(axis=1, how="all")

        self.sku_df = sheet_to_df("SKU_Parameters")
        self.dem_df = sheet_to_df("Demand")
        self.pl_df  = sheet_to_df("Production_Lines")
        self.ah_df  = sheet_to_df("Available_Hours")
        self.hc_df  = sheet_to_df("Handling_Costs")
        self.ss_df  = sheet_to_df("Safety_Stock")
        self.cp_df  = sheet_to_df("Cost_Parameters")
        self.net_df = sheet_to_df("Network")
        wb.close()

        attr_map = {
            "SKU_Parameters": "sku_df", "Demand": "dem_df",
            "Production_Lines": "pl_df", "Available_Hours": "ah_df",
            "Handling_Costs": "hc_df", "Safety_Stock": "ss_df",
            "Cost_Parameters": "cp_df", "Network": "net_df",
        }
        errors = []
        for sheet_name, req_cols in self.REQUIRED.items():
            df = getattr(self, attr_map[sheet_name])
            missing_cols = [c for c in req_cols if c not in df.columns]
            if missing_cols:
                errors.append(f"  Sheet '{sheet_name}': missing columns {missing_cols}"
                              f"\n    Found: {list(df.columns)}")
            if df.empty:
                errors.append(f"  Sheet '{sheet_name}': no data rows found.")
        if errors:
            raise SystemExit("\n  ERROR: Input validation failed:\n" +
                             "\n".join(errors))

        row_counts = {k: len(getattr(self, v)) for k, v in attr_map.items()
                      if k != "Cost_Parameters"}
        print(f"   Rows loaded: {row_counts}")
        self._parse()

    @staticmethod
    def _fix_banner_df(df, expected_cols):
        if all(c in df.columns for c in expected_cols):
            return df
        for i in range(min(10, len(df))):
            row_vals = [str(v).strip() for v in df.iloc[i].tolist()
                        if v is not None and str(v).strip() != ""]
            if sum(1 for c in expected_cols if c in row_vals) >= max(2, len(expected_cols) // 2):
                new_df = df.iloc[i + 1:].copy()
                new_df.columns = [str(v).strip() if v is not None else f"_c{ci}"
                                   for ci, v in enumerate(df.iloc[i].tolist())]
                return new_df.reset_index(drop=True).dropna(how="all")
        return df

    def _parse(self):
        for attr, cols in [
            ("cp_df",  ["Parameter", "Value"]),
            ("dem_df", ["Customer", "SKU"]),
            ("pl_df",  ["Plant", "Line", "SKU", "Month", "Line_SKU_active"]),
            ("ah_df",  ["Plant", "Line"]),
            ("hc_df",  ["Node", "SKU", "Handling_cost_per_case_INR"]),
            ("ss_df",  ["Node", "SKU", "Safety_stock_days"]),
            ("net_df", ["Source", "Destination", "SKU", "Month",
                        "Transport_cost_per_case_INR", "Lane_open"]),
            ("sku_df", ["SKU", "Revenue_per_case_INR"]),
        ]:
            setattr(self, attr, self._fix_banner_df(getattr(self, attr), cols))

        def drop_notes(df):
            if df.empty: return df
            mask = df[df.columns[0]].astype(str).str.strip().str.startswith(("↑", "►", "#", "--"))
            return df[~mask].reset_index(drop=True)

        for attr in ("cp_df","dem_df","pl_df","ah_df","hc_df","ss_df","net_df","sku_df"):
            setattr(self, attr, drop_notes(getattr(self, attr)))

        def to_num(df, cols):
            for c in cols:
                if c in df.columns:
                    df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)
            return df

        self.cp_df  = to_num(self.cp_df,  ["Value"])
        self.sku_df = to_num(self.sku_df, ["Revenue_per_case_INR"])
        self.pl_df  = to_num(self.pl_df,  ["Line_SKU_active", "Batch_size_cases",
                                            "Run_rate_cases_per_hr", "Standard_batch_run_hrs",
                                            "Min_run_hrs", "CIP_hrs",
                                            "Production_cost_per_case_INR",
                                            "RMPM_Cost_per_case_INR", "Changeover_cost_INR"])
        self.hc_df  = to_num(self.hc_df,  ["Handling_cost_per_case_INR",
                                            "Inventory_holding_cost_per_case_per_month_INR"])
        self.ss_df  = to_num(self.ss_df,  ["Safety_stock_days"])
        self.net_df = to_num(self.net_df, ["Transport_cost_per_case_INR", "Lane_open"])

        def strip_str(df, cols):
            for c in cols:
                if c in df.columns:
                    df[c] = df[c].astype(str).str.strip().replace("nan", "")
            return df

        self.pl_df  = strip_str(self.pl_df,  ["Plant", "Line", "SKU", "Month"])
        self.ah_df  = strip_str(self.ah_df,  ["Plant", "Line"])
        self.net_df = strip_str(self.net_df, ["Source", "Destination", "SKU", "Month"])
        self.dem_df = strip_str(self.dem_df, ["Customer", "SKU"])
        self.hc_df  = strip_str(self.hc_df,  ["Node", "SKU"])
        self.ss_df  = strip_str(self.ss_df,  ["Node", "SKU"])
        self.sku_df = strip_str(self.sku_df, ["SKU"])

        sheet_map = {
            "SKU_Parameters":  ("sku_df",  ["SKU", "Revenue_per_case_INR"]),
            "Demand":          ("dem_df",  ["Customer", "SKU"]),
            "Production_Lines":("pl_df",   ["Plant", "Line", "SKU", "Month",
                                            "Line_SKU_active", "Batch_size_cases",
                                            "Run_rate_cases_per_hr", "Min_run_hrs",
                                            "CIP_hrs", "Production_cost_per_case_INR",
                                            "RMPM_Cost_per_case_INR"]),
            "Available_Hours": ("ah_df",   ["Plant", "Line"]),
            "Handling_Costs":  ("hc_df",   ["Node", "SKU", "Handling_cost_per_case_INR",
                                            "Inventory_holding_cost_per_case_per_month_INR"]),
            "Safety_Stock":    ("ss_df",   ["Node", "SKU", "Safety_stock_days"]),
            "Cost_Parameters": ("cp_df",   ["Parameter", "Value"]),
            "Network":         ("net_df",  ["Source", "Destination", "SKU", "Month",
                                            "Transport_cost_per_case_INR", "Lane_open"]),
        }
        errors = []
        for sheet, (attr, req_cols) in sheet_map.items():
            df = getattr(self, attr)
            missing = [c for c in req_cols if c not in df.columns]
            if missing:
                errors.append(f"  Sheet '{sheet}': missing columns {missing}"
                              f"\n    Columns found: {list(df.columns)[:10]}")
        if errors:
            raise SystemExit("\n  INPUT VALIDATION FAILED\n" + "\n".join(errors))

        cp = dict(zip(self.cp_df["Parameter"].astype(str).str.strip(), self.cp_df["Value"]))
        self.stockout_p = float(cp.get("Stockout_penalty_per_case", 150))
        self.max_pre    = int(  cp.get("Max_prebuild_months",        2))

        id_cols        = ["Customer", "SKU"]
        self.months    = [c for c in self.dem_df.columns if c not in id_cols]
        self.T         = len(self.months)
        self.month_idx = {m: t for t, m in enumerate(self.months)}

        self.skus = self.dem_df["SKU"].unique().tolist()
        sp = self.sku_df.set_index("SKU")
        self.rev = {s: float(sp.loc[s, "Revenue_per_case_INR"])
                    for s in self.skus if s in sp.index}

        all_nodes = set(self.net_df["Source"].tolist() + self.net_df["Destination"].tolist())
        self.node_type = {}
        for n in all_nodes:
            nl = n.lower()
            if   nl.startswith("plant"): self.node_type[n] = "plant"
            elif nl.startswith("cfa"):   self.node_type[n] = "cfa"
            elif nl.startswith("cust"):  self.node_type[n] = "customer"
            else:                        self.node_type[n] = "cfa"

        self.plants    = sorted([n for n, t in self.node_type.items() if t == "plant"])
        self.cfas      = sorted([n for n, t in self.node_type.items() if t == "cfa"])
        self.customers = self.dem_df["Customer"].unique().tolist()
        self.inv_nodes = self.plants + self.cfas

        self.ln_map = {p: self.pl_df[self.pl_df["Plant"] == p]["Line"].unique().tolist()
                       for p in self.plants}
        if "Month" in self.pl_df.columns:
            self._pl = self.pl_df.set_index(["Plant", "Line", "SKU", "Month"])
            self._pl_has_month = True
        else:
            self._pl = self.pl_df.set_index(["Plant", "Line", "SKU"])
            self._pl_has_month = False
        self._ah = self.ah_df.set_index(["Plant", "Line"])

        hc = self.hc_df.set_index(["Node", "SKU"])
        self._hdl = {}; self._hold = {}
        for node in self.inv_nodes:
            for s in self.skus:
                try:
                    self._hdl[(node, s)]  = float(hc.loc[(node, s), "Handling_cost_per_case_INR"])
                    self._hold[(node, s)] = float(hc.loc[(node, s), "Inventory_holding_cost_per_case_per_month_INR"])
                except KeyError:
                    self._hdl[(node, s)] = 8.0; self._hold[(node, s)] = 3.0

        ss = self.ss_df.set_index(["Node", "SKU"])
        self._ss = {}
        for node in self.inv_nodes:
            for s in self.skus:
                try:    self._ss[(node, s)] = int(ss.loc[(node, s), "Safety_stock_days"])
                except: self._ss[(node, s)] = 7

        self._dem = {}
        for _, row in self.dem_df.iterrows():
            for m in self.months:
                self._dem[(row["Customer"], row["SKU"], m)] = float(row.get(m, 0))

        self.lane_cost = {}
        for _, row in self.net_df.iterrows():
            if int(row.get("Lane_open", 1)) != 1:
                continue
            src = str(row["Source"]).strip()
            dst = str(row["Destination"]).strip()
            sku = str(row["SKU"]).strip()
            m   = str(row["Month"]).strip()
            cost = float(row["Transport_cost_per_case_INR"])
            if m not in self.month_idx:
                continue
            t   = self.month_idx[m]
            key = (src, dst, sku, t)
            if key not in self.lane_cost or cost < self.lane_cost[key]:
                self.lane_cost[key] = cost

        self.lanes = list(self.lane_cost.items())

        print(f"   Months    : {self.months}")
        print(f"   SKUs      : {self.skus}")
        print(f"   Plants    : {self.plants}")
        print(f"   CFAs      : {self.cfas}")
        print(f"   Customers : {self.customers}")
        print(f"   Open lanes: {len(self.lanes)}  |  "
              f"Closed lane-periods: "
              f"{len(self.net_df[self.net_df.get('Lane_open', pd.Series(1, index=self.net_df.index)) == 0])}")

    def _pl_get(self, p, ln, s, col, default=0.0, month=None):
        try:
            if self._pl_has_month and month is not None:
                return float(self._pl.loc[(p, ln, s, month), col])
            elif self._pl_has_month:
                sub = self._pl.loc[(p, ln, s)]
                return float(sub[col].iloc[0]) if hasattr(sub, "iloc") else float(sub[col])
            else:
                return float(self._pl.loc[(p, ln, s), col])
        except (KeyError, TypeError, IndexError):
            return default

    def active(self, p, ln, s, month=None):
        return int(round(self._pl_get(p, ln, s, "Line_SKU_active", 1, month=month))) == 1

    def batch_size(self, p, ln, s):  return self._pl_get(p, ln, s, "Batch_size_cases")
    def run_rate(self, p, ln, s):    return self._pl_get(p, ln, s, "Run_rate_cases_per_hr", 100.0)
    def cip(self, p, ln, s):         return self._pl_get(p, ln, s, "CIP_hrs")
    def prod_cost(self, p, ln, s):   return self._pl_get(p, ln, s, "Production_cost_per_case_INR", 10)
    def rmpm(self, p, ln, s):        return self._pl_get(p, ln, s, "RMPM_Cost_per_case_INR", 0)
    def chgov(self, p, ln, s):       return self._pl_get(p, ln, s, "Changeover_cost_INR")

    def std_run_hrs(self, p, ln, s):
        stored = self._pl_get(p, ln, s, "Standard_batch_run_hrs", 0.0)
        if stored > 0: return stored
        rr = self.run_rate(p, ln, s)
        return self.batch_size(p, ln, s) / rr if rr > 0 else 0.0

    def min_run_hrs(self, p, ln, s):
        return self._pl_get(p, ln, s, "Min_run_hrs", self.std_run_hrs(p, ln, s))

    def min_batches(self, p, ln, s):
        rr = self.run_rate(p, ln, s); bs = self.batch_size(p, ln, s)
        mr = self.min_run_hrs(p, ln, s)
        if rr <= 0 or bs <= 0: return 1
        return max(1, math.ceil(rr * mr / bs))

    def hrs_per_batch(self, p, ln, s):
        rr = self.run_rate(p, ln, s)
        return self.batch_size(p, ln, s) / rr + self.cip(p, ln, s) if rr > 0 else 0.0

    def avh(self, p, ln, m):
        try:   return float(self._ah.loc[(p, ln), m])
        except: return 0.0

    def dem(self, cust, s, m):  return self._dem.get((cust, s, m), 0.0)
    def hdl(self, node, s):     return self._hdl.get((node, s), 8.0)
    def hold(self, node, s):    return self._hold.get((node, s), 3.0)
    def ss_days(self, node, s): return self._ss.get((node, s), 7)


# ── MILP Solver ───────────────────────────────────────────────────────────────
class MILPSolver:
    """
    Variables
      bat[p,ln,s,t]     INTEGER  – batches produced
      set[p,ln,s,t]     BINARY   – line setup indicator
      flow[src,dst,s,t] CONT     – cases moved on open lane
      inv[node,s,t]     CONT     – closing inventory
      unmet[cust,s,t]   CONT     – unmet demand

    Objective: minimise -(revenue - all costs)  i.e. maximise profit
    """
    def __init__(self, data: DataLoader):
        self.d = data
        self.M = data.months
        self.S = data.skus
        self.P = data.plants
        self.T = data.T
        self._build_idx()

    def _build_idx(self):
        d = self.d; idx = {}; i = 0

        self.bat_keys = []
        for p in self.P:
            for ln in d.ln_map[p]:
                for s in self.S:
                    if d.batch_size(p, ln, s) <= 0: continue
                    for t in range(self.T):
                        if not d.active(p, ln, s, month=d.months[t]): continue
                        k = ("bat", p, ln, s, t); idx[k] = i; i += 1
                        self.bat_keys.append(k)

        self.set_keys = []
        for k in self.bat_keys:
            _, p, ln, s, t = k
            sk = ("set", p, ln, s, t); idx[sk] = i; i += 1
            self.set_keys.append(sk)
        self.n_int = i

        self.flow_keys = []
        for (src, dst, s, t), _ in d.lanes:
            k = ("flow", src, dst, s, t); idx[k] = i; i += 1
            self.flow_keys.append(k)

        self.inv_keys = []
        for node in d.inv_nodes:
            for s in self.S:
                for t in range(self.T):
                    k = ("inv", node, s, t); idx[k] = i; i += 1
                    self.inv_keys.append(k)

        self.unm_keys = []
        for cust in d.customers:
            for s in self.S:
                for t in range(self.T):
                    k = ("unm", cust, s, t); idx[k] = i; i += 1
                    self.unm_keys.append(k)

        self.idx = idx; self.N = i
        print(f"\n  MILP variable counts:")
        print(f"   Batch integers  : {len(self.bat_keys)}")
        print(f"   Setup binaries  : {len(self.set_keys)}")
        print(f"   Flow variables  : {len(self.flow_keys)}")
        print(f"   Inventory       : {len(self.inv_keys)}")
        print(f"   Unmet demand    : {len(self.unm_keys)}")
        print(f"   Total           : {self.N}")

    def _objective(self):
        d = self.d; idx = self.idx
        c = np.zeros(self.N)

        for k in self.bat_keys:
            _, p, ln, s, t = k
            c[idx[k]] = (d.prod_cost(p, ln, s) + d.rmpm(p, ln, s)) * d.batch_size(p, ln, s)

        for k in self.set_keys:
            _, p, ln, s, t = k
            c[idx[k]] = d.chgov(p, ln, s)

        for k in self.flow_keys:
            _, src, dst, s, t = k
            tc  = d.lane_cost[(src, dst, s, t)]
            rev = d.rev.get(s, 0) if d.node_type.get(dst) == "customer" else 0.0
            hdl = d.hdl(dst, s)   if d.node_type.get(dst) in ("plant", "cfa") else 0.0
            c[idx[k]] = tc + hdl - rev

        for k in self.inv_keys:
            _, node, s, t = k
            c[idx[k]] = d.hold(node, s)

        for k in self.unm_keys:
            c[idx[k]] = d.stockout_p

        return c

    def _constraints(self):
        d = self.d; idx = self.idx; N = self.N
        rA = []; cA = []; vA = []; lbs = []; ubs = []; row = 0

        def add(ents, lo, hi):
            nonlocal row
            for col, val in ents:
                rA.append(row); cA.append(col); vA.append(val)
            lbs.append(lo); ubs.append(hi); row += 1

        # C1 – inventory balance at plants and CFAs
        for node in d.inv_nodes:
            for s in self.S:
                for t in range(self.T):
                    ents = []
                    if t > 0:
                        ents.append((idx[("inv", node, s, t - 1)], 1.0))
                    if node in self.P:
                        for ln in d.ln_map[node]:
                            bk = ("bat", node, ln, s, t)
                            if bk in idx:
                                ents.append((idx[bk], d.batch_size(node, ln, s)))
                    for k in self.flow_keys:
                        _, src, dst, fs, ft = k
                        if dst == node and fs == s and ft == t:
                            ents.append((idx[k], 1.0))
                    for k in self.flow_keys:
                        _, src, dst, fs, ft = k
                        if src == node and fs == s and ft == t:
                            ents.append((idx[k], -1.0))
                    ents.append((idx[("inv", node, s, t)], -1.0))
                    add(ents, 0.0, 0.0)

        # C2 – demand satisfaction
        for cust in d.customers:
            for s in self.S:
                for t, m in enumerate(self.M):
                    dem_val = d.dem(cust, s, m)
                    ents = []
                    for k in self.flow_keys:
                        _, src, dst, fs, ft = k
                        if dst == cust and fs == s and ft == t:
                            ents.append((idx[k], 1.0))
                    ents.append((idx[("unm", cust, s, t)], 1.0))
                    add(ents, dem_val, dem_val)

        # C3 – line capacity
        for p in self.P:
            for ln in d.ln_map[p]:
                for t, m in enumerate(self.M):
                    avail = d.avh(p, ln, m)
                    ents = []
                    for s in self.S:
                        bk = ("bat", p, ln, s, t)
                        if bk not in idx: continue
                        if d.batch_size(p, ln, s) > 0:
                            ents.append((idx[bk], d.hrs_per_batch(p, ln, s)))
                    if ents: add(ents, -np.inf, avail)

        # C4 – setup linking with minimum-run lower bound
        # If setup=1 the line must run at least min_batches (derived from Min_run_hrs)
        BIG_M = 500
        for bk, sk in zip(self.bat_keys, self.set_keys):
            _, p, ln, s, t = bk
            mb = d.min_batches(p, ln, s)
            add([(idx[bk], 1.0),  (idx[sk], -BIG_M)],       -np.inf, 0.0)
            add([(idx[bk], -1.0), (idx[sk], float(mb))], -np.inf, 0.0)

        # C5 – safety stock floor at every inventory node
        n_inv = max(len(d.inv_nodes), 1)
        for node in d.inv_nodes:
            for s in self.S:
                for t, m in enumerate(self.M):
                    daily  = sum(d.dem(c2, s, m) for c2 in d.customers) / 30.0 / n_inv
                    safety = daily * d.ss_days(node, s)
                    add([(idx[("inv", node, s, t)], 1.0)], safety, np.inf)

        # C6 – prebuild cap (max_pre months of forward demand)
        for p in self.P:
            for s in self.S:
                for t in range(self.T):
                    fut = sum(
                        sum(d.dem(c2, s, self.M[tt]) for c2 in d.customers)
                        for tt in range(t, min(t + d.max_pre + 1, self.T))
                    )
                    add([(idx[("inv", p, s, t)], 1.0)], -np.inf, fut)

        A  = sp.csc_matrix((vA, (rA, cA)), shape=(row, N))
        lb = np.array(lbs, dtype=float)
        ub = np.array(ubs, dtype=float)
        print(f"   Constraints     : {row}")
        return A, lb, ub

    def _integ(self):
        ig = np.zeros(self.N)
        for k in self.bat_keys: ig[self.idx[k]] = 1
        for k in self.set_keys: ig[self.idx[k]] = 1
        return ig

    def _bounds(self):
        lb = np.zeros(self.N); ub = np.full(self.N, np.inf)
        for k in self.set_keys: ub[self.idx[k]] = 1.0
        return lb, ub

    def solve(self):
        print("\n  Building MILP matrices...")
        t0    = time.time()
        c_obj = self._objective()
        A, lb_c, ub_c = self._constraints()
        ig         = self._integ()
        lb_v, ub_v = self._bounds()

        print("  Calling HiGHS MILP solver (timeout=300s, gap=2%)...")
        res = milp(
            c           = c_obj,
            constraints = LinearConstraint(A, lb_c, ub_c),
            integrality = ig,
            bounds      = Bounds(lb_v, ub_v),
            options     = {"disp": False, "time_limit": 300,
                           "mip_rel_gap": 0.02, "presolve": True},
        )
        elapsed = time.time() - t0
        sm = {0: "Optimal", 1: "Time/iter limit (partial solution)",
              2: "Infeasible", 3: "Unbounded", 4: "Infeasible or Unbounded"}
        status = sm.get(res.status, f"Unknown (code {res.status})")
        print(f"   Status  : {status}  ({elapsed:.1f}s)")

        if res.status in (2, 3, 4) or res.x is None:
            return self._diagnose_infeasibility(A, lb_c, ub_c, lb_v, ub_v,
                                                ig, c_obj, elapsed, status)
        if res.status == 1:
            print("  Warning: solver hit time limit — returning best solution found.")

        return self._extract(res.x, status, elapsed)

    def _diagnose_infeasibility(self, A, lb_c, ub_c, lb_v, ub_v,
                                ig, c_obj, elapsed, status):
        d = self.d; idx = self.idx; N = self.N

        print("\n" + "═" * 65)
        print("  INFEASIBILITY DIAGNOSTIC")
        print("  Status: " + status)
        print("═" * 65)

        findings = []

        # Step 1: LP relaxation
        print("\n  [1/3] LP relaxation...")
        lp_res = milp(
            c           = c_obj,
            constraints = LinearConstraint(A, lb_c, ub_c),
            integrality = np.zeros(N),
            bounds      = Bounds(lb_v, ub_v),
            options     = {"disp": False, "time_limit": 60, "presolve": True},
        )
        lp_feasible = (lp_res.status == 0)
        lp_status   = {0: "Feasible", 2: "Infeasible", 3: "Unbounded"}.get(
            lp_res.status, f"Code {lp_res.status}")
        print(f"     LP status: {lp_status}")

        if lp_feasible:
            root_cause = "INTEGER"
            findings.append({
                "Group": "LP Relaxation", "Status": "Feasible",
                "Violation": 0, "Worst_Constraint": "—",
                "Interpretation": ("LP is feasible. Integer batch sizes make it "
                                   "impossible to satisfy all constraints. Reduce "
                                   "batch sizes, add lines, or relax safety stock.")
            })
        else:
            root_cause = "STRUCTURAL"
            findings.append({
                "Group": "LP Relaxation", "Status": "Infeasible",
                "Violation": 0, "Worst_Constraint": "—",
                "Interpretation": ("LP is also infeasible — structural problem. "
                                   "Check network connectivity, capacity, and safety stock.")
            })

        # Step 2: constraint-group relaxation
        print("\n  [2/3] Constraint-group relaxation...")
        A_arr = A.toarray()
        n_con = A_arr.shape[0]

        labels = []
        for node in d.inv_nodes:
            for s in self.S:
                for t in range(self.T):
                    labels.append(("C1_INV_BAL", f"Inv balance {node}|{s}|t={t}"))
        for cust in d.customers:
            for s in self.S:
                for t, m in enumerate(self.M):
                    labels.append(("C2_DEMAND", f"Demand {cust}|{s}|{m}"))
        for p in self.P:
            for ln in d.ln_map[p]:
                for t, m in enumerate(self.M):
                    labels.append(("C3_CAPACITY", f"Capacity {p}|{ln}|{m}"))
        for bk in self.bat_keys:
            _, p, ln, s, t = bk
            labels.append(("C4_SETUP", f"Setup UB {p}|{ln}|{s}|t={t}"))
            labels.append(("C4_SETUP", f"Setup LB {p}|{ln}|{s}|t={t}"))
        for node in d.inv_nodes:
            for s in self.S:
                for t, m in enumerate(self.M):
                    labels.append(("C5_SAFETY", f"Safety {node}|{s}|{m}"))
        for p in self.P:
            for s in self.S:
                for t in range(self.T):
                    labels.append(("C6_PREBUILD", f"Prebuild {p}|{s}|t={t}"))
        while len(labels) < n_con:
            labels.append(("OTHER", f"Row {len(labels)}"))

        groups = ["C1_INV_BAL", "C2_DEMAND", "C3_CAPACITY", "C5_SAFETY", "C6_PREBUILD"]
        group_names = {
            "C1_INV_BAL":  "Inventory Balance (C1)",
            "C2_DEMAND":   "Demand Satisfaction (C2)",
            "C3_CAPACITY": "Line Capacity (C3)",
            "C5_SAFETY":   "Safety Stock (C5)",
            "C6_PREBUILD": "Prebuild Cap (C6)",
        }
        group_interp = {
            "C1_INV_BAL":  "Flow in/out of a node cannot balance. Check missing inbound lanes or zero batch sizes.",
            "C2_DEMAND":   "Customer demand cannot be satisfied. Check for missing or closed lanes to customers.",
            "C3_CAPACITY": "Insufficient production hours to cover demand + safety stock. Check shutdowns and CIP times.",
            "C5_SAFETY":   "Safety-stock floors conflict with capacity. Reduce safety stock days or add capacity.",
            "C6_PREBUILD": "Prebuild cap too tight. Increase Max_prebuild_months in Cost_Parameters.",
        }

        group_violations = {}
        for grp in groups:
            grp_rows = [i for i, (g, _) in enumerate(labels) if g == grp and i < n_con]
            if not grp_rows: continue
            lb_test = lb_c.copy(); ub_test = ub_c.copy()
            for r in grp_rows:
                lb_test[r] = -1e30; ub_test[r] = 1e30
            test = milp(
                c           = c_obj,
                constraints = LinearConstraint(A, lb_test, ub_test),
                integrality = np.zeros(N),
                bounds      = Bounds(lb_v, ub_v),
                options     = {"disp": False, "time_limit": 30, "presolve": True},
            )
            became_feasible = (test.status == 0)
            group_violations[grp] = became_feasible
            mark = "FEASIBLE when relaxed" if became_feasible else "still infeasible"
            print(f"     {group_names[grp]:<30}  {mark}")

        # Step 3: elastic LP for violation magnitudes
        print("\n  [3/3] Elastic LP for violation magnitudes...")
        import scipy.sparse as sp2
        N2 = N + n_con
        c2 = np.zeros(N2); c2[N:] = 1.0
        A_sp = sp.csc_matrix(A)
        I_s  = sp2.eye(n_con, format="csc")
        eq_rows = np.where(lb_c == ub_c)[0]
        lb_rows = np.where((lb_c > -1e29) & (lb_c != ub_c))[0]
        ub_rows = np.where((ub_c <  1e29) & (lb_c != ub_c))[0]
        rows_el = []; lb_el = []; ub_el = []
        A_eq_top = sp2.hstack([A_sp[eq_rows, :], -I_s[eq_rows, :]])
        A_eq_bot = sp2.hstack([A_sp[eq_rows, :],  I_s[eq_rows, :]])
        for i, r in enumerate(eq_rows):
            rhs = lb_c[r]
            rows_el.append(A_eq_top[i:i+1, :]); lb_el.append(-1e30); ub_el.append(rhs)
            rows_el.append(A_eq_bot[i:i+1, :]); lb_el.append(rhs);   ub_el.append(1e30)
        if len(lb_rows):
            A_lb = sp2.hstack([A_sp[lb_rows, :], I_s[lb_rows, :]])
            for i, r in enumerate(lb_rows):
                rows_el.append(A_lb[i:i+1, :]); lb_el.append(lb_c[r]); ub_el.append(1e30)
        if len(ub_rows):
            A_ub2 = sp2.hstack([A_sp[ub_rows, :], -I_s[ub_rows, :]])
            for i, r in enumerate(ub_rows):
                rows_el.append(A_ub2[i:i+1, :]); lb_el.append(-1e30); ub_el.append(ub_c[r])
        A_elastic  = sp2.vstack(rows_el, format="csc")
        lb_elastic = np.array(lb_el); ub_elastic = np.array(ub_el)
        bounds_el  = Bounds(
            np.concatenate([lb_v, np.zeros(n_con)]),
            np.concatenate([ub_v, np.full(n_con, 1e30)])
        )
        el_res = milp(
            c           = c2,
            constraints = LinearConstraint(A_elastic, lb_elastic, ub_elastic),
            integrality = np.zeros(N2),
            bounds      = bounds_el,
            options     = {"disp": False, "time_limit": 60, "presolve": True},
        )
        per_row_viol = {}
        if el_res.x is not None:
            slacks = el_res.x[N:]
            for i, (lbl_grp, lbl_desc) in enumerate(labels):
                if i < n_con and slacks[i] > 0.5:
                    per_row_viol.setdefault(lbl_grp, []).append((slacks[i], lbl_desc))

        for grp in groups:
            viols = per_row_viol.get(grp, [])
            relaxing_helps = group_violations.get(grp, False)
            total_viol = sum(v for v, _ in viols)
            worst_viol, worst_desc = max(viols, key=lambda x: x[0]) if viols else (0, "—")
            findings.append({
                "Group":           group_names[grp],
                "Relaxing_Helps":  "YES" if relaxing_helps else "NO",
                "Total_Violation": round(total_viol, 1),
                "Worst_Constraint":worst_desc,
                "Worst_Violation": round(worst_viol, 1),
                "Interpretation":  group_interp.get(grp, "—") if (viols or relaxing_helps)
                                   else "No violation detected.",
            })

        print("\n  INFEASIBILITY SUMMARY")
        print(f"  Root cause: {root_cause}")
        for f2 in findings[1:]:
            if f2["Relaxing_Helps"] == "YES" or f2["Total_Violation"] > 0:
                print(f"\n  {f2['Group']}")
                print(f"    Relaxing helps: {f2['Relaxing_Helps']}")
                print(f"    Violation     : {f2['Total_Violation']:,.1f}")
                print(f"    Worst row     : {f2['Worst_Constraint']}")
                print(f"    → {f2['Interpretation']}")
        print("═" * 65)

        return dict(
            infeasible = True,
            status     = status,
            elapsed    = elapsed,
            root_cause = root_cause,
            findings   = pd.DataFrame(findings),
            cap_check  = self._capacity_check(),
            lane_check = self._lane_connectivity_check(),
        )

    def _capacity_check(self):
        d = self.d; rows = []
        for t, m in enumerate(self.M):
            for s in self.S:
                total_cap = sum(
                    d.batch_size(p, ln, s) *
                    (d.avh(p, ln, m) / max(d.hrs_per_batch(p, ln, s), 0.01))
                    for p in self.P for ln in d.ln_map[p]
                    if d.active(p, ln, s, month=m) and d.batch_size(p, ln, s) > 0
                )
                total_dem = sum(d.dem(c2, s, m) for c2 in d.customers)
                n_inv     = max(len(d.inv_nodes), 1)
                total_ss  = sum(
                    (sum(d.dem(c2, s, m) for c2 in d.customers) / 30.0 / n_inv)
                    * d.ss_days(nd, s) for nd in d.inv_nodes
                )
                gap = total_cap - total_dem - total_ss
                rows.append({
                    "Month": m, "SKU": s,
                    "Total_Capacity_cases":   round(total_cap),
                    "Total_Demand_cases":     round(total_dem),
                    "Total_SafetyStock_cases":round(total_ss),
                    "Net_Gap_cases":          round(gap),
                    "Status": "OK" if gap >= 0 else "SHORTAGE",
                })
        return pd.DataFrame(rows)

    def _lane_connectivity_check(self):
        d = self.d; rows = []
        open_set = set(d.lane_cost.keys())
        for cust in d.customers:
            for s in self.S:
                for t, m in enumerate(self.M):
                    direct  = any((p, cust, s, t) in open_set for p in self.P)
                    via_cfa = any(
                        (p, cfa, s, t) in open_set and (cfa, cust, s, t) in open_set
                        for p in self.P for cfa in d.cfas
                    )
                    if not (direct or via_cfa):
                        rows.append({
                            "Customer": cust, "SKU": s, "Month": m,
                            "Direct_Lane":  "YES" if direct  else "NO",
                            "Via_CFA_Lane": "YES" if via_cfa else "NO",
                            "Reachable":    "NO — ISOLATED",
                            "Demand":       round(d.dem(cust, s, m)),
                        })
        return pd.DataFrame(rows) if rows else pd.DataFrame(
            columns=["Customer", "SKU", "Month",
                     "Direct_Lane", "Via_CFA_Lane", "Reachable", "Demand"])

    def _extract(self, x, status, elapsed):
        d = self.d; idx = self.idx
        prod_c = rmpm_c = trans_c = hdl_c = chgov_c = hold_c = stk_c = rev_t = 0.0

        for k in self.bat_keys:
            _, p, ln, s, t = k; v = x[idx[k]]
            bs = d.batch_size(p, ln, s)
            prod_c += d.prod_cost(p, ln, s) * bs * v
            rmpm_c += d.rmpm(p, ln, s)      * bs * v

        for k in self.set_keys:
            _, p, ln, s, t = k
            chgov_c += d.chgov(p, ln, s) * x[idx[k]]

        for k in self.flow_keys:
            _, src, dst, s, t = k; v = x[idx[k]]
            trans_c += d.lane_cost[(src, dst, s, t)] * v
            if d.node_type.get(dst) in ("plant", "cfa"):
                hdl_c += d.hdl(dst, s) * v
            if d.node_type.get(dst) == "customer":
                rev_t += d.rev.get(s, 0) * v

        for k in self.inv_keys:
            _, node, s, t = k
            hold_c += d.hold(node, s) * x[idx[k]]

        for k in self.unm_keys:
            stk_c += d.stockout_p * x[idx[k]]

        profit = rev_t - prod_c - rmpm_c - trans_c - hdl_c - chgov_c - hold_c - stk_c

        print(f"   Revenue             : Rs {rev_t:>16,.0f}")
        print(f"   RMPM Cost           : Rs {rmpm_c:>16,.0f}")
        print(f"   Production Cost     : Rs {prod_c:>16,.0f}")
        print(f"   Transport Cost      : Rs {trans_c:>16,.0f}")
        print(f"   Handling Cost       : Rs {hdl_c:>16,.0f}")
        print(f"   Changeover Cost     : Rs {chgov_c:>16,.0f}")
        print(f"   Holding Cost        : Rs {hold_c:>16,.0f}")
        print(f"   Stockout Penalty    : Rs {stk_c:>16,.0f}")
        print(f"   NET PROFIT          : Rs {profit:>16,.0f}")

        prod_rows=[]; flow_rows=[]; pre_rows=[]
        inv_rows=[]; ful_rows=[]; cap_rows=[]; lane_rows=[]

        for bk, sk in zip(self.bat_keys, self.set_keys):
            _, p, ln, s, t = bk; v = x[idx[bk]]
            if v < 0.5: continue
            bs        = d.batch_size(p, ln, s)
            rr        = d.run_rate(p, ln, s)
            cip_h     = d.cip(p, ln, s)
            n_batches = int(round(v))
            units     = n_batches * bs
            run_h     = (bs / rr) if rr > 0 else 0.0
            total_run = round(n_batches * run_h, 2)
            total_cip = round(n_batches * cip_h, 2)
            prod_rows.append({
                "Plant": p, "Line": ln, "SKU": s, "Month": self.M[t],
                "Batches":              n_batches,
                "Cases_Produced":       units,
                "Run_Rate_cases_per_hr":rr,
                "Std_batch_run_hrs":    round(d.std_run_hrs(p, ln, s), 2),
                "Min_run_hrs":          round(d.min_run_hrs(p, ln, s), 2),
                "Min_batches":          d.min_batches(p, ln, s),
                "Min_Run_Binding":      "YES" if n_batches == d.min_batches(p,ln,s) and d.min_batches(p,ln,s) > 1 else "NO",
                "Run_hrs_per_batch":    round(run_h, 2),
                "CIP_hrs_per_batch":    round(cip_h, 2),
                "Total_hrs_per_batch":  round(run_h + cip_h, 2),
                "Total_Run_hrs":        total_run,
                "Total_CIP_hrs":        total_cip,
                "Total_Hrs_Used":       round(total_run + total_cip, 2),
                "Production_Cost_INR":  round(d.prod_cost(p, ln, s) * units),
                "RMPM_Cost_INR":        round(d.rmpm(p, ln, s) * units),
                "Total_Mfg_Cost_INR":   round((d.prod_cost(p,ln,s) + d.rmpm(p,ln,s)) * units),
                "Handling_at_Plant_INR":round(d.hdl(p, s) * units),
                "Batch_Size_cases":     bs,
            })

        for k in self.flow_keys:
            _, src, dst, s, t = k; v = x[idx[k]]
            if v < 0.5: continue
            tc    = d.lane_cost[(src, dst, s, t)]
            rev_c = d.rev.get(s, 0) if d.node_type.get(dst) == "customer" else 0
            hdl_v = d.hdl(dst, s)   if d.node_type.get(dst) in ("plant", "cfa") else 0
            flow_rows.append({
                "Source": src, "Destination": dst,
                "Node_Type_Src": d.node_type.get(src, "?"),
                "Node_Type_Dst": d.node_type.get(dst, "?"),
                "SKU": s, "Month": self.M[t],
                "Cases_Moved":               round(v),
                "Transport_Cost_per_Case_INR":round(tc, 2),
                "Transport_Cost_Total_INR":  round(tc * v),
                "Handling_Cost_INR":         round(hdl_v * v),
                "Revenue_INR":               round(rev_c * v),
                "Net_Contribution_INR":      round((rev_c - tc - hdl_v) * v),
            })

        for k in self.inv_keys:
            _, node, s, t = k; v = x[idx[k]]
            if v > 0.5 and node in d.plants:
                pre_rows.append({
                    "Plant": node, "SKU": s, "Month_End": self.M[t],
                    "Prebuild_Cases":    round(v),
                    "Earmarked_For":     self.M[min(t + 1, self.T - 1)],
                    "Holding_Cost_INR":  round(d.hold(node, s) * v),
                })
            inv_rows.append({
                "Node": node, "Node_Type": d.node_type.get(node, "?"),
                "SKU": s, "Month": self.M[t],
                "Closing_Inventory":          round(v),
                "Safety_Stock_Days":          d.ss_days(node, s),
                "Holding_Cost_per_Case_INR":  round(d.hold(node, s), 2),
                "Holding_Cost_Total_INR":     round(d.hold(node, s) * v),
            })

        for cust in d.customers:
            for s in self.S:
                for t, m in enumerate(self.M):
                    dem_val = d.dem(cust, s, m)
                    um  = x[idx[("unm", cust, s, t)]]
                    fl  = max(0.0, dem_val - um)
                    ful_rows.append({
                        "Customer": cust, "SKU": s, "Month": m,
                        "Demand":                round(dem_val),
                        "Fulfilled":             round(fl),
                        "Unmet":                 round(um),
                        "Fulfilment_pct":        round(100 * fl / dem_val, 1) if dem_val > 0 else 100.0,
                        "Revenue_Realised_INR":  round(d.rev.get(s, 0) * fl),
                        "Revenue_Lost_INR":      round(d.rev.get(s, 0) * um),
                    })

        for p in self.P:
            for ln in d.ln_map[p]:
                for t, m in enumerate(self.M):
                    avail    = d.avh(p, ln, m)
                    used_run = 0.0; used_cip = 0.0
                    for s in self.S:
                        bk = ("bat", p, ln, s, t)
                        if bk not in idx: continue
                        n = x[idx[bk]]
                        if n < 1e-6: continue
                        rr  = d.run_rate(p, ln, s)
                        bs  = d.batch_size(p, ln, s)
                        used_run += n * (bs / rr if rr > 0 else 0.0)
                        used_cip += n * d.cip(p, ln, s)
                    used_total = used_run + used_cip
                    cap_rows.append({
                        "Plant": p, "Line": ln, "Month": m,
                        "Available_hrs":  round(avail, 1),
                        "Run_hrs_used":   round(used_run, 1),
                        "CIP_hrs_used":   round(used_cip, 1),
                        "Total_hrs_used": round(used_total, 1),
                        "Slack_hrs":      round(avail - used_total, 1),
                        "Utilisation_pct":round(100 * used_total / avail, 1) if avail > 0 else 0.0,
                        "Run_pct":        round(100 * used_run  / avail, 1) if avail > 0 else 0.0,
                        "CIP_pct":        round(100 * used_cip  / avail, 1) if avail > 0 else 0.0,
                        "Status": ("SHUTDOWN" if avail == 0
                                   else ("OVER" if used_total > avail * 1.001 else "OK")),
                    })

        for (src, dst, s, t), cost in d.lane_cost.items():
            v = x[idx[("flow", src, dst, s, t)]] if ("flow", src, dst, s, t) in idx else 0
            lane_rows.append({
                "Source": src, "Destination": dst, "SKU": s, "Month": self.M[t],
                "Lane_Status":          "OPEN",
                "Cases_Moved":          round(v),
                "Transport_Cost_INR":   round(cost * v),
                "Cost_per_Case_INR":    round(cost, 2),
            })
        for _, row in d.net_df.iterrows():
            if int(row.get("Lane_open", 1)) == 0:
                m = str(row["Month"]).strip()
                if m in d.month_idx:
                    lane_rows.append({
                        "Source": str(row["Source"]), "Destination": str(row["Destination"]),
                        "SKU": str(row["SKU"]), "Month": m,
                        "Lane_Status":        "CLOSED",
                        "Cases_Moved":        0,
                        "Transport_Cost_INR": 0,
                        "Cost_per_Case_INR":  round(float(row["Transport_cost_per_case_INR"]), 2),
                    })

        elig_rows = []
        for p in self.P:
            for ln in d.ln_map[p]:
                for s in self.S:
                    bs = d.batch_size(p, ln, s)
                    rr = d.run_rate(p, ln, s)
                    for t, m in enumerate(self.M):
                        act = d.active(p, ln, s, month=m)
                        elig_rows.append({
                            "Plant": p, "Line": ln, "SKU": s, "Month": m,
                            "Active":               1 if act else 0,
                            "Status":               "ACTIVE" if act else "BLOCKED",
                            "Batch_size_cases":     bs if act else "—",
                            "Run_rate_cases_per_hr":rr if act else "—",
                            "Std_batch_run_hrs":    round(d.std_run_hrs(p, ln, s), 2) if act else "—",
                            "Min_run_hrs":          round(d.min_run_hrs(p, ln, s), 2) if act else "—",
                            "Min_batches":          d.min_batches(p, ln, s)            if act else "—",
                            "CIP_hrs":              d.cip(p, ln, s)                    if act else "—",
                            "Prod_cost_per_case":   d.prod_cost(p, ln, s)              if act else "—",
                            "RMPM_per_case":        d.rmpm(p, ln, s)                   if act else "—",
                            "Changeover_INR":       d.chgov(p, ln, s)                  if act else "—",
                        })

        return dict(
            status=status, elapsed=elapsed,
            revenue=rev_t, prod_cost=prod_c, rmpm_cost=rmpm_c,
            trans_cost=trans_c, hdl_cost=hdl_c,
            chgov_cost=chgov_c, hold_cost=hold_c, stk_cost=stk_c,
            profit=profit,
            production   = pd.DataFrame(prod_rows),
            network_flow = pd.DataFrame(flow_rows),
            prebuild     = pd.DataFrame(pre_rows),
            inventory    = pd.DataFrame(inv_rows),
            fulfilment   = pd.DataFrame(ful_rows),
            capacity     = pd.DataFrame(cap_rows),
            lanes        = pd.DataFrame(lane_rows),
            eligibility  = pd.DataFrame(elig_rows),
        )


# ── Output Writer ─────────────────────────────────────────────────────────────
class OutputWriter:
    def __init__(self, res, data: DataLoader):
        self.r = res; self.d = data
        self.wb = openpyxl.Workbook(); self.wb.remove(self.wb.active)

    def _hdr(self, c, txt, bg=None, sz=10):
        c.value = txt; c.font = B(sz, C["wht"])
        c.fill  = F(bg or C["hm"]); c.alignment = CA(True); c.border = TB()

    def _cell(self, c, val, bg=None, fmt=None, bold=False):
        v2 = None if (isinstance(val, float) and np.isnan(val)) else val
        c.value = v2; c.font = B(10) if bold else R(10)
        c.alignment = CA(); c.border = TB()
        if bg:  c.fill = F(bg)
        if fmt: c.number_format = fmt

    def _wdf(self, ws, df, sr=1):
        if df is None or df.empty:
            ws.cell(sr, 1, "No data").font = B(10, C["acc"]); return
        for ci, col in enumerate(df.columns, 1):
            self._hdr(ws.cell(sr, ci), col)
        for ri, row in enumerate(df.itertuples(index=False), sr + 1):
            bg = C["gyl"] if ri % 2 == 0 else C["wht"]
            for ci, val in enumerate(row, 1):
                col = df.columns[ci - 1]; c = ws.cell(ri, ci); fmt = None
                if any(x in col for x in ["INR", "Revenue", "Cost", "Profit", "Contribution"]):
                    fmt = FMT_INR
                elif "pct" in col.lower():
                    fmt = FMT_PCT
                elif any(x in col.lower() for x in ["hrs", "cover", "days"]):
                    fmt = FMT_DEC
                self._cell(c, val, bg=bg, fmt=fmt)
                if "Fulfilment_pct" in col and isinstance(val, (int, float)):
                    c.fill = F(C["grnl"] if val >= 95 else (C["yell"] if val >= 80 else C["redl"]))
                if "Utilisation_pct" in col and isinstance(val, (int, float)):
                    c.fill = F(C["grnl"] if val <= 85 else (C["yell"] if val <= 95 else C["redl"]))
                if col == "Lane_Status":
                    c.fill = F(C["redl"] if val == "CLOSED" else C["grnl"])
                    c.font = B(10, C["red"] if val == "CLOSED" else C["grn"])
                if col == "Status" and isinstance(val, str):
                    if   val == "SHUTDOWN": c.fill = F(C["yell"]); c.font = B(10, C["yel"])
                    elif val == "OVER":     c.fill = F(C["redl"]); c.font = B(10, C["red"])
        for ci, col in enumerate(df.columns, 1):
            w = max(len(str(col)), *(len(str(v)) for v in df[col]), 8)
            ws.column_dimensions[get_column_letter(ci)].width = min(w + 4, 32)
        ws.freeze_panes = ws.cell(sr + 1, 1)

    def _sheet(self, title, df, tab, note=""):
        ws = self.wb.create_sheet(title); ws.tab_color = tab
        ws.sheet_view.showGridLines = False
        nc = max(len(df.columns) if df is not None and not df.empty else 1, 5)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nc)
        c = ws.cell(1, 1, title); c.font = B(13, C["wht"])
        c.fill = F(C["hd"]); c.alignment = CA(); ws.row_dimensions[1].height = 28
        sr = 2
        if note:
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=nc)
            c2 = ws.cell(2, 1, note); c2.font = R(9, C["hm"]); c2.alignment = LA()
            sr = 3
        self._wdf(ws, df, sr)
        return ws

    def _dashboard(self):
        ws = self.wb.create_sheet("Dashboard", 0)
        ws.sheet_view.showGridLines = False
        r = self.r; d = self.d

        ws.merge_cells("A1:N3")
        c = ws["A1"]
        c.value = "SUPPLY CHAIN NETWORK OPTIMISATION  ·  PROFIT MAXIMISATION  ·  MILP"
        c.font  = Font(name="Calibri", bold=True, size=20, color=C["wht"])
        c.fill  = F(C["hd"]); c.alignment = CA(); ws.row_dimensions[1].height = 44

        ws.merge_cells("A4:N4")
        ts = ws["A4"]
        ts.value = (f"Generated: {datetime.now().strftime('%d %b %Y  %H:%M')}  |  "
                    f"Solver: HiGHS MILP (scipy)  |  Status: {r['status']}  |  "
                    f"Solve time: {r['elapsed']:.1f}s  |  "
                    f"{len(d.plants)} plants · {len(d.cfas)} CFAs · "
                    f"{len(d.customers)} customers  |  "
                    f"{len(d.lanes)} open lane-SKU-periods")
        ts.font = B(9, C["wht"]); ts.fill = F(C["hm"]); ts.alignment = CA()
        ws.row_dimensions[4].height = 18

        fdf = r["fulfilment"]; cap = r["capacity"]
        ful_pct  = (100 * fdf["Fulfilled"].sum() / max(fdf["Demand"].sum(), 1)
                    if not fdf.empty else 0)
        avg_util = cap["Utilisation_pct"].mean() if not cap.empty else 0
        open_ls  = len(d.lanes)
        closed_ls = len(d.net_df[d.net_df.get("Lane_open",
                         pd.Series(1, index=d.net_df.index)) == 0])

        tiles = [
            ("NET PROFIT",           f"Rs {r['profit']:,.0f}",       C["grn"]),
            ("Revenue",              f"Rs {r['revenue']:,.0f}",       C["hm"]),
            ("RMPM Cost",            f"Rs {r['rmpm_cost']:,.0f}",     C["red"]),
            ("Production Cost",      f"Rs {r['prod_cost']:,.0f}",     C["red"]),
            ("Transport Cost",       f"Rs {r['trans_cost']:,.0f}",    C["acc"]),
            ("Handling Cost",        f"Rs {r['hdl_cost']:,.0f}",      C["ora"]),
            ("Holding Cost",         f"Rs {r['hold_cost']:,.0f}",     C["pur"]),
            ("Stockout Penalty",     f"Rs {r['stk_cost']:,.0f}",
             C["red"] if r["stk_cost"] > 0 else C["grnm"]),
            ("Demand Fulfilment",    f"{ful_pct:.1f}%",
             C["grnm"] if ful_pct >= 95 else (C["yel"] if ful_pct >= 80 else C["red"])),
            ("Avg Line Utilisation", f"{avg_util:.1f}%",
             C["grnm"] if avg_util <= 85 else (C["yel"] if avg_util <= 95 else C["red"])),
            ("Open Lanes",           f"{open_ls:,}",                  C["tbl"]),
            ("Closed Lane-Periods",  f"{closed_ls:,}",
             C["red"] if closed_ls > 0 else C["grnm"]),
        ]
        for ii, (lbl, val, col) in enumerate(tiles):
            cs = (ii % 4) * 3 + 1; rs = 6 + (ii // 4) * 3
            ws.merge_cells(start_row=rs, start_column=cs, end_row=rs, end_column=cs+2)
            cv = ws.cell(rs, cs, val)
            cv.font = Font(name="Calibri", bold=True, size=13, color=C["wht"])
            cv.fill = F(col); cv.alignment = CA(); cv.border = KB()
            ws.merge_cells(start_row=rs+1, start_column=cs, end_row=rs+1, end_column=cs+2)
            cl = ws.cell(rs+1, cs, lbl)
            cl.font = B(9, C["wht"]); cl.fill = F(col); cl.alignment = CA()
            ws.row_dimensions[rs].height = 32; ws.row_dimensions[rs+1].height = 16

        tr = 20
        ws.merge_cells(f"A{tr}:L{tr}")
        ws.cell(tr, 1, "MONTHLY P&L SNAPSHOT").font = B(12, C["hd"])
        ws.cell(tr, 1).fill = F(C["hl"]); ws.cell(tr, 1).alignment = CA()

        mhdr = ["Month", "Demand", "Fulfilled", "Unmet", "Fulfil%",
                "Revenue Rs", "RMPM Rs", "Prod Rs", "Transport Rs",
                "Handling Rs", "Holding Rs", "Net Profit Rs"]
        for ci, h in enumerate(mhdr, 1):
            self._hdr(ws.cell(tr+1, ci), h, bg=C["hd"], sz=9)

        nf = r["network_flow"]; pm = r["production"]
        for ti, m in enumerate(d.months):
            ri  = tr + 2 + ti
            fm  = fdf[fdf["Month"] == m]  if not fdf.empty else pd.DataFrame()
            nfm = nf[nf["Month"] == m]    if not nf.empty  else pd.DataFrame()
            ppm = pm[pm["Month"] == m]    if not pm.empty  else pd.DataFrame()
            inv_m = r["inventory"]
            invm  = inv_m[inv_m["Month"] == m] if not inv_m.empty else pd.DataFrame()

            dem   = fm["Demand"].sum()    if not fm.empty else 0
            fl    = fm["Fulfilled"].sum() if not fm.empty else 0
            um    = fm["Unmet"].sum()     if not fm.empty else 0
            pct   = 100 * fl / dem if dem > 0 else 100
            rev_m = nfm["Revenue_INR"].sum()              if not nfm.empty else 0
            rp_m  = ppm["RMPM_Cost_INR"].sum()            if not ppm.empty else 0
            pc_m  = ppm["Production_Cost_INR"].sum()      if not ppm.empty else 0
            tc_m  = nfm["Transport_Cost_Total_INR"].sum() if not nfm.empty else 0
            hc_m  = nfm["Handling_Cost_INR"].sum()        if not nfm.empty else 0
            hld_m = invm["Holding_Cost_Total_INR"].sum()  if not invm.empty else 0
            np_m  = rev_m - rp_m - pc_m - tc_m - hc_m - hld_m

            bg = C["gyl"] if ti % 2 == 0 else C["wht"]
            vals = [m, int(dem), int(fl), int(um), round(pct, 1),
                    int(rev_m), int(rp_m), int(pc_m), int(tc_m),
                    int(hc_m), int(hld_m), int(np_m)]
            fmts = [None, FMT_NUM, FMT_NUM, FMT_NUM, FMT_PCT,
                    FMT_INR, FMT_INR, FMT_INR, FMT_INR, FMT_INR, FMT_INR, FMT_INR]
            for ci, (v, fmt2) in enumerate(zip(vals, fmts), 1):
                c = ws.cell(ri, ci, v)
                c.font = R(9); c.alignment = CA(); c.border = TB(); c.fill = F(bg)
                if fmt2: c.number_format = fmt2
                if ci == 5:  c.fill = F(C["grnl"] if pct >= 95 else (C["yell"] if pct >= 80 else C["redl"]))
                if ci == 12: c.fill = F(C["grnl"] if np_m >= 0 else C["redl"])

        for col in range(1, 15):
            ws.column_dimensions[get_column_letter(col)].width = 16

    def _profit_bridge(self):
        ws = self.wb.create_sheet("Profit Bridge", 1)
        ws.tab_color = C["grn"]; ws.sheet_view.showGridLines = False
        r = self.r

        ws.merge_cells("A1:F2")
        c = ws["A1"]; c.value = "PROFIT & LOSS BRIDGE — FULL PLANNING HORIZON"
        c.font = Font(name="Calibri", bold=True, size=15, color=C["wht"])
        c.fill = F(C["hd"]); c.alignment = CA(); ws.row_dimensions[1].height = 32

        gp    = r["revenue"] - r["rmpm_cost"] - r["prod_cost"]
        ebdit = gp - r["trans_cost"] - r["hdl_cost"]
        items = [
            ("REVENUE",                r["revenue"],     C["grnm"], False),
            ("  (–) RMPM Cost",       -r["rmpm_cost"],   C["red"],  True),
            ("  (–) Production Cost", -r["prod_cost"],   C["red"],  True),
            ("GROSS PROFIT",           gp,               C["hm"],   False),
            ("  (–) Transport Cost",  -r["trans_cost"],  C["acc"],  True),
            ("  (–) Handling Cost",   -r["hdl_cost"],    C["ora"],  True),
            ("EBITDA (excl. fixed)",   ebdit,            C["tbl"],  False),
            ("  (–) Changeover Cost", -r["chgov_cost"],  C["ora"],  True),
            ("  (–) Holding Cost",    -r["hold_cost"],   C["pur"],  True),
            ("  (–) Stockout Penalty",-r["stk_cost"],    C["red"],  True),
            ("NET PROFIT",             r["profit"],      C["hd"],   False),
        ]
        for ci, h in enumerate(["Line Item", "Amount (Rs)", "% of Revenue"], 1):
            self._hdr(ws.cell(4, ci), h, bg=C["hd"])
        rev = max(r["revenue"], 1)
        for ii, (lbl, val, col, indent) in enumerate(items):
            ri = 5 + ii; ws.row_dimensions[ri].height = 25
            lc = ws.cell(ri, 1, lbl)
            lc.font  = B(11, C["wht"]) if not indent else R(10, C["wht"])
            lc.fill  = F(col); lc.alignment = LA(); lc.border = TB()
            ac = ws.cell(ri, 2, val)
            ac.font  = B(11, C["wht"]) if not indent else R(10, C["wht"])
            ac.fill  = F(col); ac.alignment = RA(); ac.border = TB(); ac.number_format = FMT_INR
            pc = ws.cell(ri, 3, val / rev)
            pc.font  = R(10, C["wht"]); pc.fill = F(col)
            pc.alignment = CA(); pc.border = TB(); pc.number_format = "0.0%"
        for col, w in zip("ABC", [42, 20, 16]):
            ws.column_dimensions[col].width = w

    def _sheet_elig(self, title, df):
        ws = self.wb.create_sheet(title)
        ws.tab_color = "375623"; ws.sheet_view.showGridLines = False

        months_list = df["Month"].unique().tolist() if "Month" in df.columns else []
        skus_list   = df["SKU"].unique().tolist()
        lines_combo = df[["Plant", "Line"]].drop_duplicates().values.tolist()
        nc_total    = 3 + len(months_list) + 3

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nc_total)
        c = ws.cell(1, 1, title)
        c.font = B(14, C["wht"]); c.fill = F(C["hd"]); c.alignment = CA()
        ws.row_dimensions[1].height = 30

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=nc_total)
        n2 = ws.cell(2, 1, "Green = ACTIVE  |  Red = BLOCKED this period")
        n2.font = R(9, C["hm"]); n2.alignment = LA()

        current_row = 4
        for sku in skus_list:
            ws.merge_cells(start_row=current_row, start_column=1,
                           end_row=current_row, end_column=nc_total)
            sh = ws.cell(current_row, 1, f"  SKU: {sku}")
            sh.font = B(11, C["wht"]); sh.fill = F(C["tbl"]); sh.alignment = LA()
            ws.row_dimensions[current_row].height = 20
            current_row += 1

            col_hdrs = ["Plant", "Line"] + months_list + ["Active_months", "Blocked_months"]
            for ci, h in enumerate(col_hdrs, 1):
                self._hdr(ws.cell(current_row, ci), h, bg=C["hd"], sz=9)
            ws.row_dimensions[current_row].height = 18
            current_row += 1

            for ri, (plant, ln) in enumerate(lines_combo):
                bg_base = C["gyl"] if ri % 2 == 0 else C["wht"]
                for ci_id, val_id in [(1, plant), (2, ln)]:
                    c2 = ws.cell(current_row, ci_id, val_id)
                    c2.font = B(9); c2.fill = F(bg_base); c2.border = TB(); c2.alignment = CA()

                active_cnt = 0; blocked_cnt = 0; prev_status = None
                for ci, m in enumerate(months_list, 3):
                    row_data = df[(df["Plant"] == plant) & (df["Line"] == ln) &
                                  (df["SKU"] == sku) & (df["Month"] == m)]
                    if row_data.empty:
                        cell_val = "—"; is_act = None
                    else:
                        is_act   = int(row_data.iloc[0]["Active"]) == 1
                        cell_val = "●" if is_act else "✗"
                        if is_act: active_cnt  += 1
                        else:      blocked_cnt += 1

                    c2 = ws.cell(current_row, ci, cell_val)
                    if   is_act is None: c2.fill = F(bg_base)
                    elif is_act:         c2.fill = F(C["grnl"]); c2.font = B(10, C["grn"])
                    else:                c2.fill = F(C["redl"]); c2.font = B(10, C["red"])
                    c2.alignment = CA(); c2.border = TB()
                    prev_status = is_act

                ac = ws.cell(current_row, len(months_list)+3, active_cnt)
                ac.fill = F(C["grnl"]); ac.font = B(9, C["grn"]); ac.alignment = CA(); ac.border = TB()
                bc = ws.cell(current_row, len(months_list)+4, blocked_cnt)
                bc.fill = F(C["redl"] if blocked_cnt > 0 else C["wht"])
                bc.font = B(9, C["red"]) if blocked_cnt > 0 else R(9)
                bc.alignment = CA(); bc.border = TB()
                current_row += 1
            current_row += 1

        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 10
        for ci in range(3, len(months_list)+3):
            ws.column_dimensions[get_column_letter(ci)].width = 7
        ws.column_dimensions[get_column_letter(len(months_list)+3)].width = 16
        ws.column_dimensions[get_column_letter(len(months_list)+4)].width = 16
        ws.freeze_panes = "C4"

        ws2 = self.wb.create_sheet("Eligibility Detail")
        ws2.tab_color = "375623"; ws2.sheet_view.showGridLines = False
        ws2.merge_cells("A1:M1")
        ws2.cell(1, 1, "Line Eligibility — Full Detail").font = B(11, C["wht"])
        ws2.cell(1, 1).fill = F(C["hd"]); ws2.cell(1, 1).alignment = CA()
        self._wdf(ws2, df, sr=2)

    def _infeasibility_dashboard(self):
        ws = self.wb.create_sheet("Infeasibility Report", 0)
        ws.tab_color = "C00000"; ws.sheet_view.showGridLines = False
        res = self.r

        ws.merge_cells("A1:J3")
        c = ws["A1"]; c.value = "MILP INFEASIBILITY REPORT — SUPPLY CHAIN NETWORK OPTIMISATION"
        c.font = Font(name="Calibri", bold=True, size=18, color=C["wht"])
        c.fill = F(C["red"]); c.alignment = CA(); ws.row_dimensions[1].height = 42

        ws.merge_cells("A4:J4")
        ts = ws["A4"]
        ts.value = (f"Generated: {datetime.now().strftime('%d %b %Y  %H:%M')}  |  "
                    f"Status: {res['status']}  |  Root cause: {res['root_cause']}  |  "
                    f"Solve time: {res['elapsed']:.1f}s")
        ts.font = B(10, C["wht"]); ts.fill = F("7F0000"); ts.alignment = CA()
        ws.row_dimensions[4].height = 18

        ws.merge_cells("A6:J8")
        rc = ws["A6"]
        if res["root_cause"] == "INTEGER":
            rc.value = ("ROOT CAUSE: INTEGER REQUIREMENTS — LP relaxation is FEASIBLE.\n"
                        "Batch sizes cannot be rounded to integers while satisfying all constraints.\n"
                        "Reduce batch sizes, add lines, or relax safety-stock / prebuild targets.")
        else:
            rc.value = ("ROOT CAUSE: STRUCTURAL INFEASIBILITY — LP relaxation is also INFEASIBLE.\n"
                        "The problem cannot be solved even with continuous quantities.\n"
                        "Check network connectivity, total capacity, and safety-stock levels.")
        rc.font = B(12, C["wht"])
        rc.fill = F("7F0000" if res["root_cause"] == "STRUCTURAL" else C["ora"])
        rc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        rc.border = KB()
        for r_idx in (6, 7, 8): ws.row_dimensions[r_idx].height = 22

        tr = 10
        ws.cell(tr, 1, "CONSTRAINT GROUP ANALYSIS").font = B(12, C["red"])
        ws.merge_cells(f"A{tr}:J{tr}")
        self._wdf(ws, res["findings"], tr+1)

        rr = tr + len(res["findings"]) + 4
        ws.merge_cells(f"A{rr}:J{rr}")
        ws.cell(rr, 1, "RECOMMENDED ACTIONS").font = B(12, C["hd"])
        recs = [
            "1. NETWORK:    Every Customer × SKU × Month needs ≥1 open lane. See 'Isolated Customers' sheet.",
            "2. CAPACITY:   Σ(max cases/month) ≥ Σ(demand + safety stock). See 'Capacity Check' sheet.",
            "3. SAFETY SS:  Reduce Safety_Stock_days if SS requirement exceeds production capacity.",
            "4. SHUTDOWNS:  Lines with Available_hrs=0 produce nothing — check vs peak demand months.",
            "5. BATCH SIZE: If INTEGER root cause, reduce batch sizes for finer scheduling granularity.",
            "6. PREBUILD:   Increase Max_prebuild_months to allow production further ahead.",
        ]
        for ii, rec in enumerate(recs):
            ri = rr + 1 + ii
            ws.merge_cells(f"A{ri}:J{ri}")
            c2 = ws.cell(ri, 1, rec)
            c2.font = R(10, C["hd"] if ii % 2 == 0 else "000000")
            c2.fill = F(C["gyl"] if ii % 2 == 0 else C["wht"])
            c2.alignment = LA(); c2.border = TB()

        for col in range(1, 11):
            ws.column_dimensions[get_column_letter(col)].width = 18
        ws.column_dimensions["D"].width = 40
        ws.column_dimensions["E"].width = 50

    def write(self, out):
        r = self.r
        if r.get("infeasible"):
            print("\n  Writing infeasibility report...")
            self._infeasibility_dashboard()
            self._sheet("Infeasibility Findings", r["findings"], "C00000",
                        "Groups marked YES are driving infeasibility")
            self._sheet("Capacity Check", r["cap_check"], "ED7D31",
                        "SHORTAGE = supply cannot cover demand + safety stock")
            lc = r["lane_check"]
            if not lc.empty:
                self._sheet("Isolated Customers", lc, "7030A0",
                            "No reachable open lane for this Customer × SKU × Month")
            else:
                ws = self.wb.create_sheet("Lane Connectivity"); ws.tab_color = "70AD47"
                ws.cell(1, 1, "All customers reachable in all periods.").font = B(11, C["grn"])
            self.wb.save(out)
            print(f"   Saved -> {out}")
            return out

        print("\n  Writing output workbook...")
        self._dashboard()
        self._profit_bridge()
        if not r["production"].empty:
            self._sheet("Production Schedule", r["production"], "2E75B6",
                        "Batch schedule | plant / line / SKU / month")
        if not r["network_flow"].empty:
            self._sheet("Network Flow", r["network_flow"], "ED7D31",
                        "Cases on every active lane | Plant→CFA, CFA→Customer, Plant→Customer")
        if not r["lanes"].empty:
            self._sheet("Lane Utilisation", r["lanes"], "C55A11",
                        "CLOSED lanes shown in red (Lane_open=0 in Network sheet)")
        if not r["prebuild"].empty:
            self._sheet("Prebuild Plan", r["prebuild"], "7030A0",
                        "Advance production held at plant — earmarked for future demand")
        if not r["inventory"].empty:
            self._sheet("Inventory Position", r["inventory"], "375623",
                        "Closing inventory at plants & CFAs")
        if "eligibility" in r and not r["eligibility"].empty:
            self._sheet_elig("Line Eligibility Matrix", r["eligibility"])
        if not r["fulfilment"].empty:
            self._sheet("Demand Fulfilment", r["fulfilment"], "00B050",
                        "Green ≥95%  |  Yellow 80-95%  |  Red <80%")
        if not r["capacity"].empty:
            self._sheet("Capacity Utilisation", r["capacity"], "1F7979",
                        "Green ≤85%  |  Yellow ≤95%  |  Red >95%  |  SHUTDOWN = zero hours")
        self.wb.save(out)
        print(f"   Saved -> {out}")
        return out


# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    print("=" * 70)
    print("  SUPPLY CHAIN NETWORK OPTIMISATION")
    print("  Objective : PROFIT MAXIMISATION  |  Model : MILP (scipy HiGHS)")
    print("=" * 70)

    inp = INPUT_FILE.strip()
    if not inp or inp in ("your_data.xlsx", ""):
        if len(sys.argv) > 1:
            inp = sys.argv[1].strip()
        else:
            print("\n  Edit INPUT_FILE at the top of this script, then re-run.")
            sys.exit(1)

    script_dir = os.path.dirname(os.path.abspath(__file__))
    candidates = [inp, os.path.join(script_dir, inp), os.path.join(os.getcwd(), inp)]
    resolved   = next((p for p in candidates if os.path.exists(p)), None)

    if resolved is None:
        print(f"\n  ERROR: Input file not found: {inp!r}")
        print(f"  Looked in: {candidates}")
        xlsx_nearby = [f for f in os.listdir(script_dir) if f.endswith(".xlsx")]
        if xlsx_nearby: print(f"  Excel files nearby: {xlsx_nearby}")
        sys.exit(1)

    inp = resolved
    print(f"  Input  : {inp}")

    inp_base = os.path.splitext(os.path.basename(inp))[0]
    if OUTPUT_FILE and OUTPUT_FILE.strip():
        out = OUTPUT_FILE.strip()
        os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    else:
        out = os.path.join(os.path.dirname(inp), f"{inp_base}_output.xlsx")
    print(f"  Output : {out}\n")

    data   = DataLoader(inp)
    solver = MILPSolver(data)
    res    = solver.solve()
    if res is None:
        print("  Solver returned no result."); sys.exit(1)

    OutputWriter(res, data).write(out)

    if res.get("infeasible"):
        print("\n" + "=" * 70)
        print(f"  Status     : {res['status']}")
        print(f"  Root cause : {res['root_cause']}")
        cc = res["cap_check"]
        short = cc[cc["Status"] == "SHORTAGE"] if not cc.empty else pd.DataFrame()
        if not short.empty:
            print(f"\n  Capacity shortages ({len(short)} Month×SKU):")
            for _, row in short.head(5).iterrows():
                print(f"    {row['Month']} | {row['SKU']} : gap = {row['Net_Gap_cases']:,.0f} cases")
        lc = res["lane_check"]
        if not lc.empty:
            print(f"\n  Isolated customer-SKU-months: {len(lc)}")
        print("=" * 70)
        print(f"\n  Report saved -> '{out}'\n")
        sys.exit(2)

    f = res["fulfilment"]; c2 = res["capacity"]
    print("\n" + "=" * 70)
    print("  RESULTS")
    print("=" * 70)
    print(f"  Revenue                   : Rs {res['revenue']:>16,.0f}")
    print(f"  (–) RMPM Cost             : Rs {res['rmpm_cost']:>16,.0f}")
    print(f"  (–) Production Cost       : Rs {res['prod_cost']:>16,.0f}")
    print(f"  (–) Transport Cost        : Rs {res['trans_cost']:>16,.0f}")
    print(f"  (–) Handling Cost         : Rs {res['hdl_cost']:>16,.0f}")
    print(f"  (–) Changeover Cost       : Rs {res['chgov_cost']:>16,.0f}")
    print(f"  (–) Holding Cost          : Rs {res['hold_cost']:>16,.0f}")
    print(f"  (–) Stockout Penalty      : Rs {res['stk_cost']:>16,.0f}")
    print(f"  {'─' * 46}")
    print(f"  NET PROFIT                : Rs {res['profit']:>16,.0f}")
    if res["revenue"] > 0:
        print(f"  Profit Margin             :    {100 * res['profit'] / res['revenue']:>13.1f}%")
    if not f.empty:
        p2 = 100 * f["Fulfilled"].sum() / max(f["Demand"].sum(), 1)
        print(f"  Demand Fulfilment         :    {p2:>13.1f}%")
    if not c2.empty:
        print(f"  Avg Line Utilisation      :    {c2['Utilisation_pct'].mean():>13.1f}%")
    print("=" * 70)
    print(f"\n  Open '{out}' for the complete plan.\n")


if __name__ == "__main__":
    main()
